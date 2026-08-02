"""Paramètres / Settings - clés API globales, prompts (par owner)."""
from __future__ import annotations
import io
import json
import uuid
from datetime import datetime
from typing import Annotated, Any

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from ..crypto import encrypt, decrypt
from ..database import get_db
from ..db_models import (
    CleanupConfig,
    Project,
    ProjectMemberRole,
    Prompt,
    User,
    UserCredential,
    UserRole,
)
from ..dependencies import get_current_user, require_owner, check_project_access
from ..midjourney_settings import (
    DEFAULT_GRID_WAIT_SECONDS,
    POST_UPSCALE_WAIT_SECONDS,
    UPSCALE_GAP_SECONDS,
    clamp_grid_wait,
)
from ..models import CredentialSet, CredentialOut, PromptOut, PromptsUpdate
from ..services.prompts import DEFAULT_PROMPTS

router = APIRouter(prefix="/api/settings", tags=["settings"])

VALID_KEY_TYPES = {
    "openai",
    "discord_auth",
    "discord_app_id",
    "discord_guild",
    "discord_channel",
    "mj_version",
    "mj_id",
    "google_sa_json",
}


def _mask(value: str) -> str:
    if len(value) <= 8:
        return "****"
    return value[:4] + "****" + value[-4:]


async def _resolve_owner_id(user: User) -> uuid.UUID:
    """Return the effective owner_id for credentials/prompts lookup."""
    if user.role == UserRole.owner:
        return user.id
    if user.created_by_owner_id:
        return user.created_by_owner_id
    raise HTTPException(status_code=403, detail="No owner associated with this account")


@router.get("/credentials", response_model=list[CredentialOut])
async def list_user_credentials(
    user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    owner_id = await _resolve_owner_id(user)
    result = await db.execute(
        select(UserCredential).where(UserCredential.user_id == owner_id)
    )
    creds = result.scalars().all()
    out = []
    for c in creds:
        try:
            plain = decrypt(c.encrypted_value)
        except Exception:
            plain = "****"
        out.append(CredentialOut(key_type=c.key_type, masked_value=_mask(plain), updated_at=c.updated_at))
    return out


@router.put("/credentials", response_model=list[CredentialOut])
async def set_user_credentials(
    body: list[CredentialSet],
    user: Annotated[User, Depends(require_owner)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    for item in body:
        if item.key_type not in VALID_KEY_TYPES:
            raise HTTPException(status_code=400, detail=f"Invalid key_type: {item.key_type}")
        result = await db.execute(
            select(UserCredential).where(
                UserCredential.user_id == user.id,
                UserCredential.key_type == item.key_type,
            )
        )
        cred = result.scalar_one_or_none()
        value = item.value.strip()
        if not value:
            # Empty submission means "remove this key" — never store an encrypted empty string
            # which would cause jobs to silently skip Midjourney with no error shown.
            if cred:
                await db.delete(cred)
            continue
        enc = encrypt(value)
        if cred:
            cred.encrypted_value = enc
        else:
            cred = UserCredential(
                user_id=user.id,
                key_type=item.key_type,
                encrypted_value=enc,
            )
            db.add(cred)

    await db.commit()
    result = await db.execute(
        select(UserCredential).where(UserCredential.user_id == user.id)
    )
    creds = result.scalars().all()
    return [
        CredentialOut(
            key_type=c.key_type,
            masked_value=_mask(decrypt(c.encrypted_value)) if c.encrypted_value else "****",
            updated_at=c.updated_at,
        )
        for c in creds
    ]


@router.get("/prompts", response_model=list[PromptOut])
async def list_prompts(
    user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
    project_id: uuid.UUID = Query(...),
):
    # Read access is allowed to any project member/admin/owner.
    await check_project_access(project_id, user, db)

    project_row = await db.execute(select(Project).where(Project.id == project_id))
    project = project_row.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    result = await db.execute(
        select(Prompt).where(Prompt.owner_id == project.owner_id, Prompt.project_id == project_id)
    )
    rows = result.scalars().all()
    out = {r.key: PromptOut(key=r.key, value=r.value, description=r.description or "") for r in rows}
    for key, data in DEFAULT_PROMPTS.items():
        if key not in out:
            out[key] = PromptOut(key=key, value=data["value"], description=data.get("description", ""))
    return list(out.values())


@router.delete("/prompts", status_code=204)
async def reset_prompts(
    user: Annotated[User, Depends(require_owner)],
    db: Annotated[AsyncSession, Depends(get_db)],
    project_id: uuid.UUID = Query(...),
    keys: list[str] | None = Query(default=None),
):
    """Delete selected project overrides, or every override when keys are omitted."""
    await check_project_access(
        project_id, user, db, require_roles=[ProjectMemberRole.admin]
    )
    selected_keys = set(keys) if keys is not None else None
    if selected_keys is not None:
        invalid_keys = sorted(selected_keys - DEFAULT_PROMPTS.keys())
        if invalid_keys:
            raise HTTPException(
                status_code=400,
                detail=f"Invalid prompt key: {invalid_keys[0]}",
            )
        if not selected_keys:
            return

    query = select(Prompt).where(
        Prompt.owner_id == user.id,
        Prompt.project_id == project_id,
    )
    if selected_keys is not None:
        query = query.where(Prompt.key.in_(selected_keys))
    result = await db.execute(query)
    for row in result.scalars().all():
        # Keep this guard even with the SQL filter so a future query refactor
        # cannot accidentally reset unrelated project prompts.
        if selected_keys is None or row.key in selected_keys:
            await db.delete(row)
    await db.commit()


@router.delete("/prompts/all", status_code=204)
async def reset_all_prompts(
    user: Annotated[User, Depends(require_owner)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Delete ALL custom prompt rows for every owner/project system-wide so defaults are used."""
    result = await db.execute(select(Prompt))
    for row in result.scalars().all():
        await db.delete(row)
    await db.commit()


@router.put("/prompts", response_model=list[PromptOut])
async def update_prompts(
    body: PromptsUpdate,
    user: Annotated[User, Depends(require_owner)],
    db: Annotated[AsyncSession, Depends(get_db)],
    project_id: uuid.UUID = Query(...),
):
    await check_project_access(
        project_id, user, db, require_roles=[ProjectMemberRole.admin]
    )
    for key, value in body.prompts.items():
        if key not in DEFAULT_PROMPTS:
            raise HTTPException(status_code=400, detail=f"Invalid prompt key: {key}")
        result = await db.execute(
            select(Prompt).where(
                Prompt.owner_id == user.id,
                Prompt.project_id == project_id,
                Prompt.key == key,
            )
        )
        row = result.scalar_one_or_none()
        if row:
            row.value = value
        else:
            db.add(Prompt(
                owner_id=user.id,
                project_id=project_id,
                key=key,
                value=value,
                description=DEFAULT_PROMPTS[key].get("description", ""),
            ))
    await db.commit()
    result = await db.execute(
        select(Prompt).where(Prompt.owner_id == user.id, Prompt.project_id == project_id)
    )
    rows = result.scalars().all()
    out = {r.key: PromptOut(key=r.key, value=r.value, description=r.description or "") for r in rows}
    for key, data in DEFAULT_PROMPTS.items():
        if key not in out:
            out[key] = PromptOut(key=key, value=data["value"], description=data.get("description", ""))
    return list(out.values())


# ── Pinterest Boards ─────────────────────────────────────────────────────────

_DEFAULT_BOARDS = DEFAULT_PROMPTS["pinterest_boards_list"]["value"].splitlines()


@router.post("/boards/import", response_model=dict)
async def import_boards_excel(
    user: Annotated[User, Depends(require_owner)],
    file: UploadFile = File(...),
):
    """Parse an Excel file (column A = board name, row 1 = header) and return newline-separated boards."""
    _MAX_EXCEL_BYTES = 5 * 1024 * 1024  # 5 MB
    content = await file.read()
    if len(content) > _MAX_EXCEL_BYTES:
        raise HTTPException(status_code=413, detail="File exceeds the 5 MB size limit")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content))
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid Excel file")
    ws = wb.active
    boards = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        cell = row[0] if row else None
        if cell:
            name = str(cell).strip()
            if name:
                boards.append(name)
    if not boards:
        raise HTTPException(status_code=400, detail="No board names found in column A (starting from row 2)")
    return {"boards": "\n".join(boards)}


@router.get("/boards/template")
async def get_boards_template(
    user: Annotated[User, Depends(require_owner)],
):
    """Return an Excel template pre-filled with the default Pinterest boards."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Boards"
    ws.column_dimensions["A"].width = 40
    ws.append(["Board Name"])
    for board in _DEFAULT_BOARDS:
        ws.append([board])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=pinterest_boards_template.xlsx"},
    )


# ── Midjourney: configurable grid wait only (upscale gap & download wait fixed in code) ──

class MidjourneyTimersOut(BaseModel):
    grid_wait_seconds: int
    upscale_gap_seconds: int = UPSCALE_GAP_SECONDS
    post_upscale_wait_seconds: int = POST_UPSCALE_WAIT_SECONDS


class MidjourneyGridWaitUpdate(BaseModel):
    grid_wait_seconds: int


def _parse_grid_wait_json(raw: str | None) -> int:
    if not raw:
        return DEFAULT_GRID_WAIT_SECONDS
    try:
        j = json.loads(raw)
        return clamp_grid_wait(int(j.get("grid_wait_seconds", DEFAULT_GRID_WAIT_SECONDS)))
    except (ValueError, TypeError, json.JSONDecodeError):
        return DEFAULT_GRID_WAIT_SECONDS


@router.get("/midjourney-timers", response_model=MidjourneyTimersOut)
async def get_midjourney_timers(
    user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    owner_id = await _resolve_owner_id(user)
    result = await db.execute(select(User).where(User.id == owner_id))
    owner = result.scalar_one_or_none()
    grid = _parse_grid_wait_json(owner.mj_timer_settings if owner else None)
    return MidjourneyTimersOut(
        grid_wait_seconds=grid,
        upscale_gap_seconds=UPSCALE_GAP_SECONDS,
        post_upscale_wait_seconds=POST_UPSCALE_WAIT_SECONDS,
    )


@router.put("/midjourney-timers", response_model=MidjourneyTimersOut)
async def set_midjourney_grid_wait(
    body: MidjourneyGridWaitUpdate,
    user: Annotated[User, Depends(require_owner)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    g = clamp_grid_wait(body.grid_wait_seconds)
    user.mj_timer_settings = json.dumps({"grid_wait_seconds": g})
    await db.commit()
    return MidjourneyTimersOut(
        grid_wait_seconds=g,
        upscale_gap_seconds=UPSCALE_GAP_SECONDS,
        post_upscale_wait_seconds=POST_UPSCALE_WAIT_SECONDS,
    )


# ── Cleanup Config (per owner) ───────────────────────────────────────────────


class CleanupConfigOut(BaseModel):
    enabled: bool
    interval_days: int
    last_run_at: datetime | None = None

    class Config:
        from_attributes = True


class CleanupConfigUpdate(BaseModel):
    enabled: bool
    interval_days: int = Field(ge=1, le=3650)


async def _get_or_create_cleanup_config(db: AsyncSession, owner_id: uuid.UUID) -> CleanupConfig:
    row = await db.execute(select(CleanupConfig).where(CleanupConfig.owner_id == owner_id))
    cfg = row.scalar_one_or_none()
    if not cfg:
        cfg = CleanupConfig(owner_id=owner_id, enabled=False, interval_days=7)
        db.add(cfg)
        await db.commit()
        await db.refresh(cfg)
    return cfg


@router.get("/cleanup-config", response_model=CleanupConfigOut)
async def get_cleanup_config(
    user: Annotated[User, Depends(require_owner)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    cfg = await _get_or_create_cleanup_config(db, user.id)
    return CleanupConfigOut(enabled=cfg.enabled, interval_days=cfg.interval_days, last_run_at=cfg.last_run_at)


@router.put("/cleanup-config", response_model=CleanupConfigOut)
async def update_cleanup_config(
    body: CleanupConfigUpdate,
    user: Annotated[User, Depends(require_owner)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    cfg = await _get_or_create_cleanup_config(db, user.id)
    cfg.enabled = body.enabled
    cfg.interval_days = max(1, body.interval_days)
    await db.commit()
    await db.refresh(cfg)
    return CleanupConfigOut(enabled=cfg.enabled, interval_days=cfg.interval_days, last_run_at=cfg.last_run_at)


@router.post("/cleanup-config/run-now", response_model=dict)
async def run_cleanup_now(
    user: Annotated[User, Depends(require_owner)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Immediately delete all published recipes belonging to this owner and their local images."""
    from ..services.image_retention_scheduler import run_full_published_cleanup
    result = await run_full_published_cleanup(owner_id=user.id)
    cfg = await _get_or_create_cleanup_config(db, user.id)
    from datetime import datetime, timezone
    cfg.last_run_at = datetime.now(timezone.utc)
    await db.commit()
    return result


# ── Custom Fonts (per user) ──────────────────────────────────────────────────


class FontsUpdate(BaseModel):
    fonts: list[str]


@router.get("/fonts", response_model=list[str])
async def get_fonts(
    user: Annotated[User, Depends(get_current_user)],
):
    if not user.custom_fonts:
        return []
    try:
        return json.loads(user.custom_fonts)
    except (json.JSONDecodeError, TypeError):
        return []


@router.put("/fonts", response_model=list[str])
async def set_fonts(
    body: FontsUpdate,
    user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    unique = list(dict.fromkeys(f.strip() for f in body.fonts if f.strip()))
    user.custom_fonts = json.dumps(unique)
    await db.commit()
    return unique


class ReusablePinElement(BaseModel):
    id: str = Field(min_length=1, max_length=120)
    name: str = Field(min_length=1, max_length=160)
    kind: str = Field(min_length=1, max_length=40)
    payload: dict[str, Any] = Field(default_factory=dict)
    created_at: str | None = None


class ReusablePinElementsUpdate(BaseModel):
    elements: list[ReusablePinElement]


@router.get("/pin-elements", response_model=list[ReusablePinElement])
async def get_pin_elements(
    user: Annotated[User, Depends(get_current_user)],
):
    if not user.custom_pin_elements:
        return []
    try:
        parsed = json.loads(user.custom_pin_elements)
        if not isinstance(parsed, list):
            return []
        out: list[ReusablePinElement] = []
        for item in parsed:
            try:
                out.append(ReusablePinElement.model_validate(item))
            except Exception:
                continue
        return out
    except (json.JSONDecodeError, TypeError):
        return []


@router.put("/pin-elements", response_model=list[ReusablePinElement])
async def set_pin_elements(
    body: ReusablePinElementsUpdate,
    user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    # Keep insertion order while removing duplicate ids; cap to avoid unbounded growth.
    seen: set[str] = set()
    cleaned: list[ReusablePinElement] = []
    for item in body.elements:
        element_id = item.id.strip()
        if not element_id or element_id in seen:
            continue
        seen.add(element_id)
        cleaned.append(
            ReusablePinElement(
                id=element_id,
                name=item.name.strip()[:160] or "Element",
                kind=item.kind.strip()[:40] or "unknown",
                payload=item.payload or {},
                created_at=item.created_at,
            )
        )
        if len(cleaned) >= 300:
            break

    user.custom_pin_elements = json.dumps([item.model_dump() for item in cleaned])
    await db.commit()
    return cleaned
